try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False
    print("Pandas not available, export functionality will be limited")
import os
import tempfile
import logging
from typing import List, Dict, Any
from datetime import datetime

logger = logging.getLogger(__name__)


class ExportService:
    """Service for exporting BOM data to various formats"""
    
    def __init__(self):
        self.temp_dir = tempfile.gettempdir()
        self.export_columns = [
            'item_number', 'part_number', 'description', 'quantity',
            'unit', 'material', 'supplier', 'cost'
        ]
    
    def export_to_csv(self, bom_items: List[Any], original_filename: str) -> str:
        """
        Export BOM items to CSV format
        
        Args:
            bom_items: List of BOMItem model instances
            original_filename: Original document filename
            
        Returns:
            Path to the generated CSV file
        """
        try:
            # Convert BOM items to DataFrame
            df = self._items_to_dataframe(bom_items)
            
            # Generate output filename
            base_name = os.path.splitext(original_filename)[0]
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"{base_name}_BOM_{timestamp}.csv"
            output_path = os.path.join(self.temp_dir, output_filename)
            
            # Export to CSV
            df.to_csv(output_path, index=False, encoding='utf-8')
            
            logger.info(f"BOM exported to CSV: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"CSV export failed: {str(e)}")
            raise
    
    def export_to_xlsx(self, bom_items: List[Any], original_filename: str) -> str:
        """
        Export BOM items to Excel format
        
        Args:
            bom_items: List of BOMItem model instances
            original_filename: Original document filename
            
        Returns:
            Path to the generated Excel file
        """
        try:
            # Convert BOM items to DataFrame
            df = self._items_to_dataframe(bom_items)
            
            # Generate output filename
            base_name = os.path.splitext(original_filename)[0]
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"{base_name}_BOM_{timestamp}.xlsx"
            output_path = os.path.join(self.temp_dir, output_filename)
            
            # Create Excel writer with formatting
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                # Write main BOM data
                df.to_excel(writer, sheet_name='BOM', index=False)
                
                # Get workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['BOM']
                
                # Apply formatting
                self._format_excel_worksheet(workbook, worksheet, df)
                
                # Add summary sheet
                self._add_summary_sheet(writer, df, original_filename)
            
            logger.info(f"BOM exported to Excel: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Excel export failed: {str(e)}")
            raise
    
    def _items_to_dataframe(self, bom_items: List[Any]) -> Any:
        """Convert BOM items to pandas DataFrame"""
        data = []
        
        for item in bom_items:
            row = {}
            for column in self.export_columns:
                # Handle both model instances and dictionaries
                if hasattr(item, column):
                    value = getattr(item, column)
                else:
                    value = item.get(column, '') if isinstance(item, dict) else ''
                
                # Clean and format values
                if value is None:
                    value = ''
                elif isinstance(value, str):
                    value = value.strip()
                
                row[self._format_column_name(column)] = value
            
            data.append(row)
        
        if PANDAS_AVAILABLE:
            return pd.DataFrame(data)
        else:
            return data
    
    def _format_column_name(self, column: str) -> str:
        """Format column name for display"""
        name_mapping = {
            'item_number': 'Item No.',
            'part_number': 'Part Number',
            'description': 'Description',
            'quantity': 'Qty',
            'unit': 'Unit',
            'material': 'Material',
            'supplier': 'Supplier',
            'cost': 'Cost'
        }
        
        return name_mapping.get(column, column.replace('_', ' ').title())
    
    def _format_excel_worksheet(self, workbook: Any, worksheet: Any, df: Any) -> None:
        """Apply formatting to Excel worksheet"""
        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            
            # Header formatting
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            # Apply header formatting
            for col in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            # Data formatting
            data_alignment = Alignment(horizontal="left", vertical="center")
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Apply data formatting
            for row in range(1, len(df) + 2):  # +2 for header and 1-based indexing
                for col in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.alignment = data_alignment
                    cell.border = thin_border
            
            # Auto-adjust column widths
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)  # Max width of 50
                worksheet.column_dimensions[column_letter].width = adjusted_width
        
        except ImportError:
            logger.warning("openpyxl formatting not available")
        except Exception as e:
            logger.warning(f"Excel formatting failed: {str(e)}")
    
    def _add_summary_sheet(self, writer: Any, df: Any, original_filename: str) -> None:
        """Add summary sheet to Excel workbook"""
        try:
            # Create summary data
            summary_data = {
                'Document Information': [''],
                'Original File': [original_filename],
                'Export Date': [datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
                'Total Items': [len(df)],
                '': [''],
                'BOM Statistics': [''],
                'Items with Part Numbers': [len(df[df['Part Number'].notna() & (df['Part Number'] != '')])],
                'Items with Descriptions': [len(df[df['Description'].notna() & (df['Description'] != '')])],
                'Items with Quantities': [len(df[df['Qty'].notna() & (df['Qty'] != '')])],
                'Items with Costs': [len(df[df['Cost'].notna() & (df['Cost'] != '')])],
            }
            
            summary_df = pd.DataFrame(list(summary_data.items()), columns=['Metric', 'Value'])
            summary_df.to_excel(writer, sheet_name='Summary', index=False)
            
            # Format summary sheet
            summary_sheet = writer.sheets['Summary']
            
            try:
                from openpyxl.styles import Font, PatternFill
                
                # Format header rows
                for row in [1, 6]:  # Header rows
                    if row <= len(summary_df):
                        cell = summary_sheet.cell(row=row + 1, column=1)
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")
            
            except ImportError:
                pass
        
        except Exception as e:
            logger.warning(f"Summary sheet creation failed: {str(e)}")
    
    def export_validation_report(self, job_id: str, validation_results: Dict[str, Any], 
                               original_filename: str) -> str:
        """
        Export validation report with issues and confidence scores
        
        Args:
            job_id: Job identifier
            validation_results: Results from validation service
            original_filename: Original document filename
            
        Returns:
            Path to the generated validation report
        """
        try:
            # Generate output filename
            base_name = os.path.splitext(original_filename)[0]
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"{base_name}_validation_report_{timestamp}.xlsx"
            output_path = os.path.join(self.temp_dir, output_filename)
            
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                # Summary sheet
                summary_data = {
                    'Overall Confidence': [f"{validation_results['overall_confidence'] * 100:.1f}%"],
                    'Total Items': [validation_results['metadata']['total_items']],
                    'Valid Items': [validation_results['metadata']['valid_items']],
                    'Items with Errors': [validation_results['metadata']['items_with_errors']],
                    'Items with Warnings': [validation_results['metadata']['items_with_warnings']],
                    'Average Item Confidence': [f"{validation_results['validation_summary']['average_confidence'] * 100:.1f}%"]
                }
                
                summary_df = pd.DataFrame(list(summary_data.items()), columns=['Metric', 'Value'])
                summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
                # Item details sheet
                if validation_results['item_validations']:
                    items_data = []
                    for validation in validation_results['item_validations']:
                        items_data.append({
                            'Item Index': validation['item_index'],
                            'Confidence Score': f"{validation['confidence_score'] * 100:.1f}%",
                            'Errors': len(validation['errors']),
                            'Warnings': len(validation['warnings']),
                            'Status': 'Valid' if validation['confidence_score'] >= 0.8 else 'Needs Review'
                        })
                    
                    items_df = pd.DataFrame(items_data)
                    items_df.to_excel(writer, sheet_name='Item Details', index=False)
                
                # Issues sheet
                issues_data = []
                for validation in validation_results['item_validations']:
                    for error in validation['errors']:
                        issues_data.append({
                            'Item Index': validation['item_index'],
                            'Severity': 'Error',
                            'Field': error['field'],
                            'Message': error['message']
                        })
                    for warning in validation['warnings']:
                        issues_data.append({
                            'Item Index': validation['item_index'],
                            'Severity': 'Warning',
                            'Field': warning['field'],
                            'Message': warning['message']
                        })
                
                if issues_data:
                    issues_df = pd.DataFrame(issues_data)
                    issues_df.to_excel(writer, sheet_name='Issues', index=False)
            
            logger.info(f"Validation report exported: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Validation report export failed: {str(e)}")
            raise
